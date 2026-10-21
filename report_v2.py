# -*- coding: utf-8 -*-
"""Independent, engineer-style Word report; never claims to be an official inspection."""
from __future__ import annotations
import math
from pathlib import Path
from datetime import datetime


def _number(x, digits=2, suffix=''):
    if not isinstance(x, (int, float)) or isinstance(x, bool) or not math.isfinite(x):
        return '未提供'
    return f'{x:,.{digits}f}{suffix}'


def _safe(v, default='未提供'):
    return str(v).strip() if v is not None and str(v).strip() else default


def build_report(path, measured, context, evidence, analysis, *, image=None, overlay=None,
                 rag_state='live', timestamp=None):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_BREAK

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.15), Cm(2.15)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.95)

    ink = RGBColor(30, 48, 64)
    navy = RGBColor(30, 75, 100)
    muted = RGBColor(95, 111, 126)
    rust = RGBColor(148, 79, 51)

    def set_font(style, size, color=ink, bold=False):
        s = doc.styles[style]
        s.font.name = 'Microsoft YaHei'
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = color
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        return s

    normal = set_font('Normal', 9.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2
    h1 = set_font('Heading 1', 14, navy, True)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(7)
    h2 = set_font('Heading 2', 10.5, navy, True)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(5)
    set_font('Title', 23, navy, True)
    set_font('Caption', 8, muted)

    def shaded(cell, color='EDF3F7'):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), color)
        tcPr.append(shd)

    def repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        trPr.append(el)

    def table(headers, values, col_widths=None):
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = 'Table Grid'
        tb.autofit = False
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        repeat_header(tb.rows[0])
        for j, head in enumerate(headers):
            cell = tb.rows[0].cells[j]
            cell.text = str(head)
            shaded(cell)
        for i, vrow in enumerate(values):
            cells = tb.add_row().cells
            for j, val in enumerate(vrow):
                cells[j].text = str(val)
                if i % 2:
                    shaded(cells[j], 'F7FAFC')
        for row in tb.rows:
            for j, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if col_widths:
                    cell.width = Cm(col_widths[j])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(8.1)
                        run.font.name = 'Microsoft YaHei'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return tb

    def info(kvs):
        return table(['项目', '内容'], [(a, _safe(b)) for a, b in kvs], [4.2, 12.45])

    def add_p(text, *, color=None, bold=False, center=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = bold
        if color:
            r.font.color.rgb = color
        return p

    hd = sec.header.paragraphs[0]
    hd.text = 'BRIDGE / AI  ·  桥梁表观病害智能识别与辅助评估'
    hd.style = doc.styles['Caption']
    hd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ft = sec.footer.paragraphs[0]
    ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ft.add_run('技术演示报告（非检测机构出具）  ·  数值来源须核实   |   第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    ft._p.append(field)

    mode = measured.get('calibration', {}).get('mode')
    simulated = mode == 'simulation'
    scale = measured.get('calibration', {}).get('mm_per_pixel')
    badge = '假设比例尺演示' if simulated else '无物理标定 / 仅像素量化'
    now = timestamp or datetime.now().strftime('%Y-%m-%d')
    m = measured.get('pixel_metrics', {})
    phy = measured.get('physical_metrics', {})

    # COVER, only schematic fields; no copied enterprise identifiers.
    doc.add_paragraph('\n\n\n')
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('桥梁表观病害智能检测\n辅助评估报告')
    add_p('AI INSPECTION / PIXEL GEOMETRY / STANDARD RETRIEVAL', color=muted, center=True)
    doc.add_paragraph('\n')
    add_p('初赛系统原型 · 技术演示版', center=True, bold=True, color=navy)
    add_p(f'测量状态：{badge}', center=True, bold=True, color=rust)
    doc.add_paragraph('\n')
    info([
        ('案例编号', context.get('case_id')),
        ('工程项目', context.get('project_name')),
        ('检测对象', context.get('component_type')),
        ('构件编号', context.get('component_id')),
        ('检测区域', context.get('component_location')),
        ('生成日期', now),
        ('报告性质', '非正式工程检测报告，不具有现场检测证明效力'),
    ])
    doc.add_paragraph('\n')
    add_p('数据组成：现有分割量化结果 + 独立标定演示 + 规范候选检索 + AI 辅助成文', center=True, color=muted)
    add_p('说明：本示例不使用企业检测案例的图像、标识、签章及业务信息。', center=True, color=muted)

    # Page 2 structured statement and detail, mirror inspection report columns, no fake rating.
    doc.add_page_break()
    doc.add_heading('一、检测对象及数据来源', level=1)
    info([
        ('桥梁或结构类型', context.get('structure_type')),
        ('构件名称/编号', f"{_safe(context.get('component_type'))} / {_safe(context.get('component_id'))}"),
        ('构件具体部位', context.get('component_location')),
        ('病害类别/方向', f"{_safe(context.get('damage_type'))} / {_safe(context.get('crack_orientation'))}"),
        ('构件信息来源', context.get('component_type_source', '人工演示填写；未现场核实')),
        ('观测范围', f"{m.get('observed_crop_percent', 0):.4f}% 是病害相对于裁剪观察画面的占比，非构件面积比"),
    ])
    doc.add_heading('二、病害明细登记', level=1)
    if simulated:
        visible = _number(phy.get('visible_main_path_length_mm_demo') / 1000 if isinstance(phy.get('visible_main_path_length_mm_demo'),(int,float)) else None, 4, ' m*')
        width = _number(phy.get('normal_width_mean_mm_demo'), 3, ' mm*')
    else:
        visible, width = '无标定', '无标定'
    table(['编号', '构件', '病害类型', '所在位置', '数量', '可见长度', '宽度均值', '照片'], [
        ('D-001', _safe(context.get('component_id'), '未定'), '裂缝（走向待核）',
         _safe(context.get('component_location'), '未定'),
         f"{m.get('candidate_connected_regions', 0)} 候选域", visible, width, 'P-01 / P-02')
    ], [1.5, 1.8, 2.15, 3.25, 1.75, 2.45, 2.0, 1.6])
    add_p('注：* 表示人为假设比例尺换算结果，非真实测量。候选连通域数量不等于独立裂缝的人工确认数量。', color=rust)
    doc.add_heading('三、像素级测量与单位转换', level=1)
    table(['量化项目', '像素算法输出', '尺度演示值 / 可用性'], [
        ('最长可见主路径', _number(m.get('visible_main_path_length_px_approx'), 2, ' px'), _number(phy.get('visible_main_path_length_mm_demo'), 3, ' mm*') if simulated else '不可用'),
        ('修剪骨架网络长度', _number(m.get('pruned_skeleton_network_length_px_approx'), 2, ' px'), '仅供几何描述，不等同于单裂缝长度'),
        ('法线平均宽度', _number(m.get('normal_width_mean_px_approx'), 3, ' px'), _number(phy.get('normal_width_mean_mm_demo'), 3, ' mm*') if simulated else '不可用'),
        ('法线最大宽度', _number(m.get('normal_width_max_px_approx'), 3, ' px'), _number(phy.get('normal_width_max_mm_demo'), 3, ' mm*') if simulated else '不可用'),
        ('观察区病害像素面积', _number(m.get('pixel_area_px2'), 0, ' px²'), _number(phy.get('projected_crack_area_mm2_demo'), 3, ' mm²*') if simulated else '不可用'),
        ('有效法线测宽覆盖', _number(100*(m.get('normal_sample_coverage_fraction') or 0), 2, '%'), _safe(m.get('normal_sample_count'), '0') + ' 个测点'),
    ], [5.0, 5.4, 6.3])
    add_p(f'标定假设：s = {_number(scale, 4, " mm/px")}；长度 = 像素长度 × s；投影面积 = 像素面积 × s²。' if simulated else '未输入可信标尺、成像距离或其他物理尺度，不能换算毫米单位。')
    add_p('本次裂缝触及裁剪图像边界，结果描述的是可见范围，不能推断完整裂缝的实际长度。', color=rust)

    # Page 3 figures
    doc.add_page_break()
    doc.add_heading('四、表观病害照片与模型辅助图', level=1)
    imgs = [('P-01  输入观察图', image), ('P-02  病害分割叠加图', overlay)]
    for name, img in imgs:
        doc.add_heading(name, level=2)
        if img and Path(img).is_file():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img), width=Cm(12.8))
        else:
            add_p('未提供相应图像；仅保留影像编号和量化记录。', color=muted)
    add_p('图片为算法数据证据。现场桥名、梁号及距离墩台位置等工程定位信息不能从该局部图像中自动推定。', color=muted)

    # Page 4 RAG provenance and analysis
    doc.add_page_break()
    doc.add_heading('五、检索规范与依据状态', level=1)
    add_p('检索类型：' + ('实时 ChromaDB 检索' if rag_state == 'live' else ('现有检索缓存复核' if rag_state == 'file' else '未执行规范检索')))
    if evidence:
        table(['序号', '规范来源 / 标准', '条款定位', '核验状态'], [
            (str(i), _safe(e.get('standard_id')) + '\n' + _safe(e.get('file_name')),
             _safe(e.get('clause_id')), '原件待核验' if not e.get('citation_verified') else '来源标记已核验')
            for i, e in enumerate(evidence[:4], 1)
        ], [1.3, 9.6, 2.25, 3.55])
    else:
        add_p('本次未执行规范检索，未输出条款匹配或工程标度。', color=rust)
    add_p('注意：相似检索结果只是候选依据。当前文本块和 OCR 表格未经正式标准原件核对，不能作为自动生成病害标度及整桥等级的依据。', color=rust)

    doc.add_heading('六、智能辅助分析与处置建议', level=1)
    add_p('文本生成来源：' + _safe(analysis.get('source')) + '；不等同于持证检测人员意见。', color=muted)
    for lab, k in [('6.1 病害信息归纳', 'summary'), ('6.2 规范适用性分析', 'evidence_analysis'),
                   ('6.3 资料局限', 'limitations'), ('6.4 复核与处置建议', 'suggestions')]:
        doc.add_heading(lab, level=2)
        add_p(_safe(analysis.get(k), '该项未提供'))

    doc.add_heading('七、应用边界与复核清单', level=1)
    limits = [
        '物理尺度：本次毫米值仅用于换算过程演示，未依据实测标尺或经过验证的成像几何。',
        '观察范围：裂缝被 640×480 中心裁剪边界截断；照片面积比例不代表构件面积比例。',
        '结构定位：构件身份、桥跨及裂缝真实结构方向是外部工程信息，需人工核验。',
        '规范使用：条款版本、原件表格、受力与材料条件需核验；不得自动推断桥梁安全等级或 Dr。',
        '正式应用：需补充现场标定、独立尺寸验证、全面桥检及专业人员复核。',
    ]
    for item in limits:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)
    doc.add_heading('八、技术结论', level=1)
    add_p('本技术演示展示了“病害分割结果→像素量化→尺度换算接口→桥梁规范候选检索→AI 辅助分析→结构化 Word 报告”的软件链路。当前不是现场检测报告，不输出整桥技术状况评分、实际合格判定或结构安全结论。')
    doc.core_properties.author = 'Bridge AI Demo'
    doc.core_properties.title = '桥梁表观病害智能检测辅助评估报告（技术演示版）'
    doc.core_properties.subject = 'AI-assisted demonstration, not an official inspection'
    doc.save(path)
    return path


# V2.5: unify single-image and batch reports; keep original export callable.
_old_build_report = build_report
def build_report(path, measured, context, evidence, analysis, *, image=None, overlay=None, rag_state='live', timestamp=None):
    from batch_report import build_consolidated_docx
    from inspection_synthesis import findings_for_measurement
    from dynamic_rag import detected_damages
    from pathlib import Path
    img=Path(image).name if image else '单张检测图'
    row={'case_id':context.get('case_id','SINGLE-001'),'relative_path':img,
         'source_path':str(image or ''),'image_path':str(image or ''),
         'overlay_path':str(overlay or ''),'quant_json':'',
         'status':'completed','measurements':measured,
         'detected_damage_codes':detected_damages(measured),
         'engineering_context':context,
         'inspection_findings':findings_for_measurement(measured,context.get('case_id','SINGLE-001'))}
    cal=measured.get('calibration') or {}
    manifest={'root':'单图','scale_mode':cal.get('mode'),
              'demo_mm_per_px':cal.get('mm_per_pixel')}
    return build_consolidated_docx(Path(path),manifest,[row],context,evidence,analysis,rag_mode=rag_state)
